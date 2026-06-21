from docx import Document
from openpyxl import load_workbook


def extract_docx(file_path):
    doc = Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])


def extract_xlsx(file_path):
    wb = load_workbook(file_path, data_only=True)

    text = ""

    for sheet in wb.sheetnames:
        ws = wb[sheet]

        for row in ws.iter_rows(values_only=True):
            row_text = " ".join([str(c) for c in row if c])
            text += row_text + "\n"

    return text