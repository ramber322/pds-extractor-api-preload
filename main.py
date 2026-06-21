from fastapi import FastAPI, UploadFile, File, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from docx import Document
from openpyxl import load_workbook
import os
import re
import numpy as np
import joblib
from typing import List, Dict, Any, Optional
from datetime import datetime

app = FastAPI(
    title="AI-Enhanced PDS Screening API",
    version="3.0",
    description="Upload PDS and screen against LGU/CSC job requirements with relevance filtering"
)

# =========================
# CORS MIDDLEWARE
# =========================
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =========================
# LOAD MODEL (optional)
# =========================
MODEL_PATH = "model.pkl"
model = joblib.load(MODEL_PATH) if os.path.exists(MODEL_PATH) else None

from feature_engine import (
    get_education_level,
    get_experience_years,
    get_training_hours,
    get_eligibility,
    detect_eligibility_level,
    get_eligibility_match_status
)

# =========================
# JOB TITLE TO CATEGORY MAPPING
# =========================
JOB_TITLE_TO_CATEGORY = {
    "administrative": ["administrative assistant", "administrative aide", "assessment clerk", "clerk", "secretary", "admin", "office staff"],
    "finance": ["budgeting assistant", "budget", "accountant", "accounting", "cashier", "treasury", "finance"],
    "engineering": ["engineer", "civil engineer", "mechanical engineer", "electrical engineer", "engineering"],
    "construction": ["building inspector", "labor foreman", "construction", "foreman", "inspector"],
    "communications": ["information officer", "public information", "media", "communications", "journalism", "writing"],
    "agriculture": ["slaughterhouse master", "agriculturist", "veterinary", "agriculture", "veterinarian"],
    "sports": ["sports officer", "recreation", "sports"],
    "technical": ["mechanic", "technician", "technical"],
    "it": ["it officer", "information technology", "it", "tech support", "developer"],
    "planning": ["planning officer", "development officer", "planning"],
    "social_services": ["social worker", "community affairs", "social", "community"]
}

# =========================
# DEFAULT JOB REQUIREMENTS BY CATEGORY
# =========================
DEFAULT_JOB_REQUIREMENTS = {
    "administrative": {
        "position_title": "Administrative Officer",
        "education": 3,
        "experience": 1,
        "training_hours": 4,
        "eligibility_level": "professional",
        "description": "Administrative support and office management roles"
    },
    "finance": {
        "position_title": "Budgeting Assistant / Accountant",
        "education": 3,
        "experience": 1,
        "training_hours": 4,
        "eligibility_level": "professional",
        "description": "Budgeting, accounting, procurement, and financial management roles"
    },
    "engineering": {
        "position_title": "Engineer III",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "Civil, mechanical, electrical engineering and design roles"
    },
    "construction": {
        "position_title": "Building Inspector / Foreman",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "Building inspection, construction supervision, and safety roles"
    },
    "communications": {
        "position_title": "Information Officer III",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "Public information, media relations, and communications roles"
    },
    "agriculture": {
        "position_title": "Slaughterhouse Master / Agriculturist",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "Agriculture, veterinary, food safety, and inspection roles"
    },
    "sports": {
        "position_title": "Sports Officer",
        "education": 3,
        "experience": 1,
        "training_hours": 4,
        "eligibility_level": "professional",
        "description": "Sports development, recreation, and physical education roles"
    },
    "technical": {
        "position_title": "Mechanic / Technician",
        "education": 2,
        "experience": 1,
        "training_hours": 4,
        "eligibility_level": "subprofessional",
        "description": "Mechanical, automotive, and equipment maintenance roles"
    },
    "it": {
        "position_title": "Information Technology Officer",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "IT, systems, networking, and software development roles"
    },
    "planning": {
        "position_title": "Planning Officer",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "Urban planning, development, and research roles"
    },
    "social_services": {
        "position_title": "Social Worker / Community Affairs Officer",
        "education": 3,
        "experience": 2,
        "training_hours": 8,
        "eligibility_level": "professional",
        "description": "Social work, community development, and welfare roles"
    }
}

# =========================
# CLEAN TEXT
# =========================
def clean(text: Any) -> str:
    if text is None:
        return ""
    return str(text).strip()

def is_noise_row(row: List[str]) -> bool:
    text = " ".join(row).lower()
    
    noise_patterns = [
        "personal data sheet", "warning:", "read the attached guide",
        "print legibly", "cs form", "page", "signature", "date",
        "wet signature", "continue on separate sheet", "do not abbreviate",
        "indicate n/a", "tick appropriate boxes", "please indicate country",
        "write in full", "dd/mm/yyyy", "if applicable", "separate sheet",
        "name of children", "photo", "right thumbmark", "subscribed and sworn",
        "i. personal information", "ii. family background", "iii. educational background",
        "iv. civil service", "v. work", "vi. voluntary", "vii. learning",
        "viii. other information", "ix.", "x.", "pursuant to", "references",
        "government issued id", "person administering oath", "affiant exhibiting",
        "indigenous people", "magna carta", "expanded solo parents",
        "candidate in a national", "resigned from the government",
        "immigrant or permanent resident", "status of appointment", "gov't service"
    ]
    
    for pattern in noise_patterns:
        if pattern in text:
            return True
    
    if len(text) < 3:
        return True
        
    return False

# =========================
# DOCX EXTRACT
# =========================
def extract_docx(file_path: str) -> List[List[str]]:
    doc = Document(file_path)
    rows = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and len(text) > 3:
            parts = re.split(r'\t{2,}|\s{3,}', text)
            cleaned = [clean(p) for p in parts if clean(p)]
            if cleaned and not is_noise_row(cleaned):
                rows.append(cleaned)
    
    return rows

# =========================
# XLSX EXTRACT
# =========================
def extract_xlsx(file_path: str) -> List[List[str]]:
    wb = load_workbook(file_path, data_only=True)
    rows = []
    
    relevant_sheets = ["C1", "C2", "C3"]
    
    for sheet_name in relevant_sheets:
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        
        for row in ws.iter_rows(values_only=True):
            cleaned = [clean(c) for c in row if clean(c)]
            if cleaned:
                rows.append(cleaned)
    
    return rows

# =========================
# SMART PDS NORMALIZER
# =========================
def normalize_pds(rows: List[List[str]]) -> Dict[str, List[List[str]]]:
    data = {
        "education": [],
        "experience": [],
        "eligibility": [],
        "training": []
    }
    
    in_education = False
    in_experience = False
    in_eligibility = False
    in_training = False
    
    seen_education_header = False
    seen_experience_header = False
    seen_eligibility_header = False
    seen_training_header = False
    
    for row in rows:
        if not row:
            continue
            
        text = " ".join(row).lower()
        
        # === DETECT SECTION HEADERS ===
        if "educational background" in text or "iii. educational" in text:
            in_education = True
            in_experience = False
            in_eligibility = False
            in_training = False
            seen_education_header = True
            continue
            
        if "work experience" in text or "v. work" in text:
            in_experience = True
            in_education = False
            in_eligibility = False
            in_training = False
            seen_experience_header = True
            continue
            
        if "civil service" in text or "eligibility" in text or "cse" in text or "csee" in text:
            in_eligibility = True
            in_education = False
            in_experience = False
            in_training = False
            seen_eligibility_header = True
            continue
            
        if "learning and development" in text or "training programs" in text or "vii. learning" in text:
            in_training = True
            in_education = False
            in_experience = False
            in_eligibility = False
            seen_training_header = True
            continue
        
        # === STORE DATA ===
        if in_education and seen_education_header:
            if any(word in text for word in ["write in full", "period of attendance", "units earned", "level", "name of school"]):
                continue
            if len(row) >= 3 and "n/a" not in text:
                data["education"].append(row)
                
        elif in_experience and seen_experience_header:
            if any(word in text for word in ["write in full", "inclusive dates", "dd/mm/yyyy", "status of appointment"]):
                continue
            if len(row) >= 3:
                has_date = False
                for cell in row[:2]:
                    cell_str = str(cell).lower()
                    if isinstance(cell, datetime) or re.search(r'(20\d{2}|19\d{2})', cell_str) or "present" in cell_str:
                        has_date = True
                        break
                if has_date:
                    data["experience"].append(row)
                    
        elif in_eligibility and seen_eligibility_header:
            if any(word in text for word in ["rating", "date of examination", "license", "valid until", "ces/csee", "conferment"]):
                continue
            
            all_na_or_empty = True
            for cell in row:
                if cell is None:
                    continue
                cell_str = str(cell).strip().lower()
                if cell_str and cell_str not in ["n/a", "none", ""]:
                    all_na_or_empty = False
                    break
            
            if all_na_or_empty:
                continue
            
            if len(row) >= 1:
                data["eligibility"].append(row)
                
        elif in_training and seen_training_header:
            if any(word in text for word in ["write in full", "inclusive dates", "number of hours", "type of l&d"]):
                continue
            if len(row) >= 3 and "n/a" not in text:
                data["training"].append(row)
    
    return data

# =========================
# FEATURE ENGINE
# =========================
def build_features(pds: Dict[str, List[List[str]]], job_category: str = "administrative") -> Dict[str, Any]:
    education_text = " ".join([" ".join([str(cell) for cell in row]) for row in pds["education"]])
    eligibility_text = " ".join([" ".join([str(cell) for cell in row]) for row in pds["eligibility"]])
    
    education_level = get_education_level(education_text)
    years_experience = get_experience_years(pds["experience"], job_category)
    training_hours = get_training_hours(pds["training"], job_category)
    
    eligibility_has = get_eligibility(eligibility_text)
    eligibility_level = detect_eligibility_level(eligibility_text)
    
    return {
        "education_level": education_level,
        "years_experience": years_experience,
        "training_hours": training_hours,
        "eligibility": eligibility_has,
        "eligibility_level": eligibility_level
    }

# =========================
# SCORING ENGINE - HYBRID SCORING (Partial Credit + Binary)
# =========================
def compute_score(applicant: Dict[str, Any], job: Dict[str, Any]) -> Dict[str, Any]:
    breakdown = {}
    sum_of_progress = 0
    total_requirements = 0
    met_requirements = 0

    # === EDUCATION ===
    req_edu = job.get("education", 1)
    val_edu = applicant.get("education", 0)
    
    if req_edu == 0:
        match_progress = 1.0
        match_binary = 1.0
        status = "NOT REQUIRED"
    elif val_edu >= req_edu:
        match_progress = 1.0
        match_binary = 1.0
        status = "MET"
    else:
        match_progress = val_edu / req_edu
        match_binary = 0.0
        status = f"{round(match_progress * 100)}% (needs {req_edu - val_edu} more level)"
    
    breakdown["Education"] = {
        "required": req_edu,
        "actual": val_edu,
        "progress": round(match_progress, 2),
        "binary": match_binary,
        "status": status
    }
    
    if status != "NOT REQUIRED":
        total_requirements += 1
        sum_of_progress += match_progress
        if match_binary >= 1:
            met_requirements += 1

    # === EXPERIENCE ===
    req_exp = job.get("experience", 1)
    val_exp = applicant.get("experience", 0)
    
    if req_exp == 0:
        match_progress = 1.0
        match_binary = 1.0
        status = "NOT REQUIRED"
    elif val_exp >= req_exp:
        match_progress = 1.0
        match_binary = 1.0
        status = "MET"
    else:
        match_progress = val_exp / req_exp
        match_binary = 0.0
        status = f"{round(match_progress * 100)}% (needs {req_exp - val_exp} more years)"
    
    breakdown["Experience"] = {
        "required": req_exp,
        "actual": val_exp,
        "progress": round(match_progress, 2),
        "binary": match_binary,
        "status": status
    }
    
    if status != "NOT REQUIRED":
        total_requirements += 1
        sum_of_progress += match_progress
        if match_binary >= 1:
            met_requirements += 1

    # === TRAINING ===
    req_train = job.get("training_hours", 1)
    val_train = applicant.get("training_hours", 0)
    
    if req_train == 0:
        match_progress = 1.0
        match_binary = 1.0
        status = "NOT REQUIRED"
    elif val_train >= req_train:
        match_progress = 1.0
        match_binary = 1.0
        status = "MET"
    else:
        match_progress = val_train / req_train if req_train > 0 else 1.0
        match_binary = 0.0
        status = f"{round(match_progress * 100)}% (needs {req_train - val_train} more hours)"
    
    breakdown["Training Hours"] = {
        "required": req_train,
        "actual": val_train,
        "progress": round(match_progress, 2),
        "binary": match_binary,
        "status": status
    }
    
    if status != "NOT REQUIRED":
        total_requirements += 1
        sum_of_progress += match_progress
        if match_binary >= 1:
            met_requirements += 1

    # === ELIGIBILITY ===
    req_elig_level = job.get("eligibility_level", "professional")
    app_elig_level = applicant.get("eligibility_level", "unknown")
    
    if req_elig_level == "none":
        match_progress = 1.0
        match_binary = 1.0
        status = "NOT REQUIRED"
    else:
        elig_result = get_eligibility_match_status(app_elig_level, req_elig_level)
        match_progress = elig_result["score"]
        match_binary = elig_result["score"]
        status = f"{elig_result['status']} ({elig_result['message']})"
    
    breakdown["Eligibility"] = {
        "required": req_elig_level,
        "actual": app_elig_level,
        "progress": round(match_progress, 2),
        "binary": round(match_binary, 2),
        "status": status
    }
    
    if status != "NOT REQUIRED":
        total_requirements += 1
        sum_of_progress += match_progress
        if match_binary >= 1:
            met_requirements += 1

    # === CALCULATE HYBRID SCORE ===
    if total_requirements == 0:
        hybrid_score = 1.0
        hybrid_percentage = 100.0
        binary_percentage = 100.0
        requirements_met = "0/0"
        verdict = "FIT"
    else:
        # Hybrid score = sum of progress / total requirements (gives partial credit)
        hybrid_score = sum_of_progress / total_requirements
        hybrid_percentage = round(hybrid_score * 100, 2)
        
        # Binary percentage = strict requirements met
        binary_percentage = round((met_requirements / total_requirements) * 100, 2)
        requirements_met = f"{met_requirements}/{total_requirements}"
        verdict = "FIT" if met_requirements == total_requirements else "NOT FIT"
    
    return {
        "score": round(hybrid_score, 2),
        "percentage": hybrid_percentage,
        "binary_percentage": binary_percentage,
        "requirements_met": requirements_met,
        "breakdown": breakdown,
        "verdict": verdict,
        "met_requirements": met_requirements,
        "total_requirements": total_requirements
    }

# =========================
# XAI EXPLANATION
# =========================
def explain_prediction(applicant: Dict[str, Any], job: Dict[str, Any], score_result: Dict[str, Any], job_title: str = "") -> Dict[str, Any]:
    explanation = {
        "summary": "",
        "verdict_description": "",
        "job_title": job_title,
        "feature_breakdown": [],
        "recommendations": []
    }
    
    breakdown = score_result["breakdown"]
    
    for feature, data in breakdown.items():
        status = data["status"]
        actual = data["actual"]
        required = data["required"]
        progress = data.get("progress", 0)
        
        # Determine if requirement is met
        is_met = False
        
        if status == "MET":
            is_met = True
        elif "EXCEEDS" in status:
            is_met = True
        elif "MET" in status and "NOT" not in status:
            is_met = True
        elif "MATCH" in status and "NO" not in status.upper():
            is_met = True
        
        # Determine description
        if "NOT REQUIRED" in status:
            desc = "Not required"
        elif is_met:
            desc = "Meets requirement"
            if feature == "Eligibility" and "(" in status and ")" in status:
                match = re.search(r'\((.*?)\)', status)
                if match:
                    desc = match.group(1)
                    if desc:
                        desc = desc[0].upper() + desc[1:] if len(desc) > 1 else desc
        else:
            desc = "Does not meet requirement"
            if feature == "Eligibility":
                explanation["recommendations"].append(
                    f"Obtain the required eligibility: {required}"
                )
            elif feature == "Education":
                gap = required - actual
                if gap > 0:
                    explanation["recommendations"].append(
                        f"Complete {int(gap)} more education level(s) (currently level {actual}, needs level {required})"
                    )
            elif feature == "Experience":
                gap = required - actual
                if gap > 0:
                    explanation["recommendations"].append(
                        f"Gain {int(gap)} more years of relevant experience"
                    )
            elif feature == "Training Hours":
                gap = required - actual
                if gap > 0:
                    explanation["recommendations"].append(
                        f"Complete {int(gap)} more hours of relevant training"
                    )
        
        explanation["feature_breakdown"].append({
            "feature": feature,
            "actual": actual,
            "required": required,
            "progress": round(progress * 100, 2),
            "status": status,
            "description": desc
        })
    
    # === SUMMARY LOGIC ===
    met_count = score_result.get("met_requirements", 0)
    total = score_result.get("total_requirements", 0)
    
    if total == 0:
        explanation["summary"] = "No specific requirements for this position. You are eligible!"
    elif met_count == total:
        explanation["summary"] = f"Applicant meets all {total} requirements! Perfect match."
    elif met_count >= total * 0.75:
        explanation["summary"] = f"Applicant meets {met_count} out of {total} requirements. Strong candidate with minor gaps."
    elif met_count > 0:
        explanation["summary"] = f"Applicant meets only {met_count} out of {total} requirements. Needs improvement."
    else:
        explanation["summary"] = f"Applicant meets NO requirements."
    
    explanation["verdict_description"] = (
        "Qualified and recommended for further consideration." if score_result["verdict"] == "FIT"
        else "Does not meet all requirements. Review recommendations above."
    )
    
    return explanation

# =========================
# ML PREDICTION
# =========================
def ml_predict(features: Dict[str, Any]) -> int:
    if not model:
        return None
    
    import warnings
    warnings.filterwarnings("ignore", category=UserWarning)
    
    X = np.array([[
        features["education_level"],
        features["years_experience"],
        features["training_hours"],
        features["eligibility"]
    ]])
    
    pred = model.predict(X)[0]
    return int(pred)

# =========================
# DETECT JOB CATEGORY
# =========================
def detect_job_category(position_title: str) -> str:
    title_lower = position_title.lower()
    
    for category, keywords in JOB_TITLE_TO_CATEGORY.items():
        for keyword in keywords:
            if keyword in title_lower:
                return category
    
    return "administrative"

# =========================
# MAIN ENDPOINT - UPLOAD PDS
# =========================
@app.post("/upload-pds")
async def upload_pds(
    file: UploadFile = File(..., description="PDS file (Excel or Word)"),
    position_title: Optional[str] = Query(None, description="Position title"),
    job_category: Optional[str] = Query(None, description="Job category"),
    education: Optional[int] = Query(None, description="Required education level (0-4)"),
    experience: Optional[int] = Query(None, description="Required years of experience"),
    training_hours: Optional[int] = Query(None, description="Required training hours"),
    eligibility_level: Optional[str] = Query(None, description="Required eligibility: 'professional', 'subprofessional', 'none'")
):
    temp_file = f"temp_{file.filename}"
    
    try:
        with open(temp_file, "wb") as f:
            f.write(await file.read())
        
        if file.filename.lower().endswith(".docx"):
            raw = extract_docx(temp_file)
        elif file.filename.lower().endswith((".xlsx", ".xls")):
            raw = extract_xlsx(temp_file)
        else:
            return JSONResponse(
                status_code=400,
                content={"error": "Unsupported file format. Please upload .docx, .xlsx, or .xls"}
            )
        
        if not raw:
            return JSONResponse(
                status_code=400,
                content={"error": "No data could be extracted from the file"}
            )
        
        if job_category:
            detected_category = job_category.lower()
        elif position_title:
            detected_category = detect_job_category(position_title)
        else:
            detected_category = "administrative"
        
        if detected_category not in DEFAULT_JOB_REQUIREMENTS:
            detected_category = "administrative"
        
        defaults = DEFAULT_JOB_REQUIREMENTS[detected_category]
        
        final_education = education if education is not None else defaults["education"]
        final_experience = experience if experience is not None else defaults["experience"]
        final_training = training_hours if training_hours is not None else defaults["training_hours"]
        final_eligibility = eligibility_level if eligibility_level is not None else defaults["eligibility_level"]
        
        pds = normalize_pds(raw)
        features = build_features(pds, detected_category)
        
        applicant = {
            "education": features["education_level"],
            "experience": features["years_experience"],
            "training_hours": features["training_hours"],
            "eligibility": features["eligibility"],
            "eligibility_level": features["eligibility_level"]
        }
        
        job = {
            "education": final_education,
            "experience": final_experience,
            "training_hours": final_training,
            "eligibility_level": final_eligibility
        }
        
        score_result = compute_score(applicant, job)
        job_title_display = position_title if position_title else defaults.get("position_title", "Position")
        explanation = explain_prediction(applicant, job, score_result, job_title_display)
        ml_result = ml_predict(features)
        
        return {
            "success": True,
            "filename": file.filename,
            "position_title": job_title_display,
            "job_category": detected_category,
            "job_description": defaults.get("description", ""),
            "applicant_features": applicant,
            "job_requirements": job,
            "fit_analysis": score_result,
            "explanation": explanation,
            "ml_prediction": ml_result,
            "raw_data_summary": {
                "education_entries": len(pds["education"]),
                "experience_entries": len(pds["experience"]),
                "eligibility_entries": len(pds["eligibility"]),
                "training_entries": len(pds["training"])
            }
        }
    
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse(
            status_code=500,
            content={"error": str(e)}
        )
    
    finally:
        if os.path.exists(temp_file):
            os.remove(temp_file)

# =========================
# GET ELIGIBILITY OPTIONS
# =========================
@app.get("/eligibility-options")
async def get_eligibility_options():
    return {
        "eligibility_options": [
            {"value": "professional", "label": "Career Service (Professional) Second Level"},
            {"value": "subprofessional", "label": "Career Service (Subprofessional) First Level"},
            {"value": "none", "label": "None Required"}
        ]
    }

# =========================
# GET JOB CATEGORIES
# =========================
@app.get("/job-categories")
async def get_job_categories():
    result = {}
    for category, data in DEFAULT_JOB_REQUIREMENTS.items():
        result[category] = {
            "position_title": data["position_title"],
            "default_requirements": {
                "education": data["education"],
                "experience": data["experience"],
                "training_hours": data["training_hours"],
                "eligibility_level": data["eligibility_level"]
            },
            "description": data["description"]
        }
    return {"categories": result}

# =========================
# HEALTH CHECK
# =========================
@app.get("/health")
async def health_check():
    return {"status": "healthy", "model_loaded": model is not None}